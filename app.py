import os
import csv
import io
from flask import Flask, render_template, request, jsonify, send_file, redirect, url_for
import database
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from urllib.parse import quote

app = Flask(__name__, static_folder='static', static_url_path='/static', template_folder='templates')

# Garantir inicialização do Banco de Dados
database.init_db()
database.seed_sample_data_if_empty()

@app.route('/')
def index():
    return redirect(url_for('admin'))

@app.route('/admin')
def admin():
    return render_template('admin.html')

@app.route('/atualizar/<token>')
def customer_portal(token):
    mensalista = database.get_mensalista_by_token(token)
    if not mensalista:
        return render_template('error.html', mensagem="Link de atualização inválido ou expirado."), 404
    return render_template('customer.html', mensalista=mensalista)

@app.route('/sucesso')
def success_page():
    return render_template('success.html')

# ==================== API ENDPOINTS ==================== #

@app.route('/api/mensalistas', methods=['GET'])
def api_get_mensalistas():
    mensalistas = database.get_all_mensalistas()
    host_url = request.host_url.rstrip('/')
    
    # Adicionar link completo de atualização em cada item
    for m in mensalistas:
        m['link_atualizacao'] = f"{host_url}/atualizar/{m['token']}"
        # Gerar link formatado do WhatsApp
        configs = database.get_configuracoes()
        msg_template = configs.get('mensagem_template', 'Olá {nome}, atualize seu cadastro: {link}')
        nome_estacionamento = configs.get('nome_estacionamento', 'Estacionamento WPS')
        
        msg_personalizada = msg_template.replace('{nome}', m['nome'])\
                                        .replace('{link}', m['link_atualizacao'])\
                                        .replace('{nome_estacionamento}', nome_estacionamento)
                                        
        m['mensagem_pronta'] = msg_personalizada
        m['whatsapp_url'] = f"https://api.whatsapp.com/send?phone={m['telefone']}&text={quote(msg_personalizada)}"
        
    return jsonify(mensalistas)

@app.route('/api/mensalistas', methods=['POST'])
def api_add_mensalista():
    data = request.json or {}
    nome = data.get('nome', '').strip()
    telefone = data.get('telefone', '').strip()
    
    if not nome or not telefone:
        return jsonify({'error': 'Nome e telefone são obrigatórios.'}), 400
        
    m_id = database.add_mensalista(nome, telefone)
    return jsonify({'success': True, 'id': m_id}), 201

@app.route('/api/mensalistas/<int:m_id>', methods=['DELETE'])
def api_delete_mensalista(m_id):
    database.delete_mensalista(m_id)
    return jsonify({'success': True})

@app.route('/api/importar-csv', methods=['POST'])
def api_import_csv():
    if 'file' not in request.files:
        return jsonify({'error': 'Nenhum arquivo enviado.'}), 400
        
    file = request.files['file']
    filename = file.filename.lower()
    
    importados = 0
    erros = 0
    
    try:
        if filename.endswith('.csv'):
            stream = io.StringIO(file.stream.read().decode("utf-8", errors="ignore"), newline=None)
            csv_input = csv.reader(stream, delimiter=';' if ';' in stream.getvalue() else ',')
            
            # Reset stream
            stream.seek(0)
            headers = next(csv_input, None)
            
            for row in csv_input:
                if len(row) >= 2:
                    nome = row[0].strip()
                    telefone = row[1].strip()
                    if nome and telefone and not nome.lower().startswith('nome'):
                        database.add_mensalista(nome, telefone)
                        importados += 1
        elif filename.endswith('.xlsx') or filename.endswith('.xls'):
            wb = openpyxl.load_workbook(file)
            sheet = wb.active
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if row and len(row) >= 2 and row[0] and row[1]:
                    nome = str(row[0]).strip()
                    telefone = str(row[1]).strip()
                    if nome and telefone and not nome.lower().startswith('nome'):
                        database.add_mensalista(nome, telefone)
                        importados += 1
        else:
            return jsonify({'error': 'Formato não suportado. Envie CSV ou Excel (.xlsx).'}), 400
            
        return jsonify({'success': True, 'importados': importados, 'erros': erros})
    except Exception as e:
        return jsonify({'error': f"Erro ao processar arquivo: {str(e)}"}), 500

@app.route('/api/configuracoes', methods=['GET', 'POST'])
def api_configuracoes():
    if request.method == 'GET':
        return jsonify(database.get_configuracoes())
    else:
        data = request.json or {}
        for k, v in data.items():
            database.set_configuracao(k, str(v))
        return jsonify({'success': True})

@app.route('/api/token/<token>', methods=['GET'])
def api_get_token(token):
    mensalista = database.get_mensalista_by_token(token)
    if not mensalista:
        return jsonify({'error': 'Token inválido'}), 404
    return jsonify(mensalista)

@app.route('/api/recadastro/<token>', methods=['POST'])
def api_post_recadastro(token):
    data = request.json or {}
    veiculos = data.get('veiculos', [])
    
    if not veiculos or not isinstance(veiculos, list):
        return jsonify({'error': 'Informe ao menos 1 veículo com Modelo e Placa.'}), 400
        
    success, msg = database.update_mensalista_veiculos(token, veiculos)
    if not success:
        return jsonify({'error': msg}), 400
        
    return jsonify({'success': True, 'mensagem': msg})

@app.route('/api/marcar-enviado/<int:m_id>', methods=['POST'])
def api_marcar_enviado(m_id):
    database.update_status_envio(m_id, 'Enviado')
    return jsonify({'success': True})

@app.route('/api/exportar-wps', methods=['GET'])
def api_exportar_wps():
    """Gera um arquivo Excel (.xlsx) formatado pronto para o sistema WPS."""
    mensalistas = database.get_all_mensalistas()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cadastro Veiculos WPS"
    
    # Estilização do cabeçalho
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Dark slate
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    headers = [
        "ID Mensalista", "Nome do Mensalista", "Telefone/WhatsApp",
        "Modelo do Veículo", "Ano", "Cor", "Placa Liberada (LPR)", "Data Recadastro"
    ]
    
    ws.append(headers)
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        
    row_count = 2
    for m in mensalistas:
        veiculos = m.get('veiculos', [])
        if veiculos:
            for v in veiculos:
                ws.append([
                    m['id'],
                    m['nome'],
                    m['telefone'],
                    v['modelo'],
                    v['ano'],
                    v['cor'],
                    v['placa'],
                    m['data_atualizacao'] or 'Pendente'
                ])
                for col_num in range(1, 9):
                    cell = ws.cell(row=row_count, column=col_num)
                    cell.border = thin_border
                    cell.alignment = align_center if col_num in [1, 3, 5, 7, 8] else align_left
                row_count += 1
        else:
            # Caso ainda não tenha atualizado os veículos
            ws.append([
                m['id'],
                m['nome'],
                m['telefone'],
                "Pendente de Recadastro",
                "-",
                "-",
                "-",
                "Pendente"
            ])
            for col_num in range(1, 9):
                cell = ws.cell(row=row_count, column=col_num)
                cell.border = thin_border
                cell.alignment = align_center if col_num in [1, 3, 5, 7, 8] else align_left
            row_count += 1
            
    # Ajustar largura das colunas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"Relatorio_Veiculos_WPS_{timestamp}.xlsx"
    )

if __name__ == '__main__':
    print("Iniciando servidor MSGWHATS na porta 5000...")
    app.run(host='0.0.0.0', port=5000, debug=True)
